"""Build a lead sheet (.xlsx) from your contacts export, its labels and any cards read.

Usage: python3 -m wxmac.wx_contacts_excel --out ~/Desktop/wechat-leads.xlsx

Reads <data dir>/contacts/all_contacts.jsonl + classified.jsonl (+ cards.jsonl). Needs openpyxl.
"""
import argparse, json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

DATA = config.DATA / "contacts"
TIER_ORDER = {"A": 0, "B": 1, "C": 2, "-": 3}
HEAD_FILL = PatternFill("solid", fgColor="DBE7FD")
CARD_FIELDS = [("region", "Region (card)"), ("signature", "Signature"), ("source", "Added via"), ("added", "Added on"),
               ("shared_groups", "Shared groups"), ("channels", "Channels"), ("weixin_id", "Weixin ID")]
BASE = [("tier", "Tier"), ("cat", "Category"), ("remark", "Your remark"), ("alias", "Nickname"),
        ("company", "Company / brand"), ("role", "Role"), ("industry", "Industry"), ("region_hint", "Region hint"),
        ("tags", "Tags"), ("why", "Why")]


def load(name: str) -> list[dict]:
    p = DATA / name
    return [json.loads(l) for l in p.open()] if p.exists() else []


def sheet(wb: Workbook, title: str, rows: list[dict], cols: list[tuple[str, str]], widths: dict):
    ws = wb.create_sheet(title)
    ws.append([c[1] for c in cols])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
    for r in rows:
        ws.append([r.get(k, "") for k, _ in cols])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (k, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 14)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args()
    contacts = load("all_contacts.jsonl")
    cls = {o["i"]: o for o in load("classified.jsonl")}
    cards = {o["i"]: o for o in load("cards.jsonl")}
    rows = []
    for i, c in enumerate(contacts):
        k = cls.get(i, {})
        r = {"no": i + 1, **c, "cat": k.get("cat", "unlabelled"), "tier": k.get("tier", ""),
             "company": k.get("company", ""), "role": k.get("role", ""), "industry": k.get("industry", ""),
             "region_hint": k.get("region", ""), "why": k.get("why", "")}
        r.update({f: cards.get(i, {}).get(f, "") for f, _ in CARD_FIELDS})
        rows.append(r)
    order = lambda r: (TIER_ORDER.get(r["tier"], 9), r["cat"], r["remark"] or r["alias"])
    widths = {"remark": 22, "alias": 22, "company": 16, "why": 26, "tags": 20, "signature": 24, "source": 16,
              "cat": 10, "tier": 6}
    cols = BASE + CARD_FIELDS + [("next", "Next step (yours)"), ("no", "Row")]

    wb = Workbook()
    wb.remove(wb.active)
    leads = sorted([r for r in rows if r["tier"] in ("A", "B")], key=order)
    groups = [("Leads A+B", leads),
              ("Customers", sorted([r for r in rows if r["cat"] == "customer"], key=order)),
              ("Investors", sorted([r for r in rows if r["cat"] == "investor"], key=order)),
              ("Peers & partners", sorted([r for r in rows if r["cat"] in ("peer", "partner")], key=order)),
              ("All contacts", sorted(rows, key=order))]

    ws = wb.create_sheet("Read me")
    by_cat, by_tier = Counter(r["cat"] for r in rows), Counter(r["tier"] for r in rows)
    for line in [
        ["WeChat contacts lead sheet"],
        ["Source", "Manage Contacts table (nickname, remark, tags) read by OCR; cards opened for the chosen tiers"],
        ["How", "Read-only: no message sent, no chat opened, unread marks untouched"],
        ["Contacts", len(rows)], ["Labelled", len(cls)], ["Cards read", len(cards)], [],
        ["Tier", "Meaning", "Count"],
        ["A", "Clearly a decision maker in a business that buys what you sell", by_tier.get("A", 0)],
        ["B", "Runs or works in a business; role unclear", by_tier.get("B", 0)],
        ["C", "Weak signal (incl. peers who may refer)", by_tier.get("C", 0)],
        ["-", "Not a lead (friends, consumer services, investors, unknown)", by_tier.get("-", 0)], [],
        ["Category", "Count"], *[[k, v] for k, v in by_cat.most_common()], [],
        ["Note", "Labels rest only on the visible text, which OCR may misread. Your own remarks are the most reliable."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 14, 70, 10
    for title, rs in groups:
        sheet(wb, f"{title} ({len(rs)})", rs, cols, widths)

    ws = wb.create_sheet("Lead mix")
    ws.append(["Industry (A+B)", "Count", "", "Region (A+B)", "Count"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
    ti = Counter(r["industry"] for r in leads if r["industry"]).most_common(40)
    tr = Counter(r["region_hint"] for r in leads if r["region_hint"]).most_common(40)
    for j in range(max(len(ti), len(tr))):
        ws.append([*(ti[j] if j < len(ti) else ("", "")), "", *(tr[j] if j < len(tr) else ("", ""))])
    ws.column_dimensions["A"].width = ws.column_dimensions["D"].width = 22

    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"saved {a.out}: {len(rows)} contacts, {len(leads)} A+B")


if __name__ == "__main__":
    main()
